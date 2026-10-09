# -*- coding: utf-8 -*-
"""
@Remark: 导入导出功能（对应 Django utils/export_import_mixin.py 的 ImportExportMixin）
支持：导出Excel（嵌套字段）、导入Excel（嵌套字段映射+图片字段）、下载导入模板
由 CustomViewSet.register 挂载为 /export_data/ /import_data/ /download_template/ 路由
"""
import logging
import os
import random
import time
from io import BytesIO
from urllib.parse import urlparse, unquote

from fastapi.responses import Response as FastAPIResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from application.settings import MEDIA_ROOT
from utils.jsonResponse import DetailResponse, ErrorResponse

logger = logging.getLogger(__name__)

IMAGE_WHITE_LIST = ['.png', '.jpg', '.jpeg', '.gif']


def is_image(urlstr) -> bool:
    """是否为图片地址"""
    if not isinstance(urlstr, str):
        return False
    try:
        if not (urlstr.startswith('http://') or urlstr.startswith('https://')):
            return False
        a = urlparse(urlstr)
        file_name = os.path.basename(a.path)
        _, file_suffix = os.path.splitext(file_name)
        return file_suffix.lower() in IMAGE_WHITE_LIST
    except Exception:
        return False


def is_number(num) -> bool:
    """是否为数字"""
    try:
        float(num)
        return True
    except (TypeError, ValueError):
        pass
    try:
        import unicodedata
        unicodedata.numeric(num)
        return True
    except (TypeError, ValueError):
        pass
    return False


def get_string_len(string, export_column_width=80):
    """获取字符串最大宽度（中文按2.1计）"""
    length = 4
    if string is None:
        return length
    if is_number(string):
        string = str(string)
    for char in string:
        length += 2.1 if ord(char) > 256 else 1
    return round(length, 1) if length <= export_column_width else export_column_width


def inserImg(sheet, Imgpath, row_index, col_index):
    """插入图片（网络图片；失败忽略写URL）"""
    try:
        import requests
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        response = requests.get(Imgpath, timeout=10)
        img = BytesIO(response.content)
        imgToInsert = XLImage(img)
        imgsize_excel = XDRPositiveSize2D(pixels_to_EMU(40), pixels_to_EMU(40))
        marker = AnchorMarker(col=col_index - 1, colOff=pixels_to_EMU(8), row=row_index - 1, rowOff=pixels_to_EMU(5))
        imgToInsert.anchor = OneCellAnchor(_from=marker, ext=imgsize_excel)
        sheet.add_image(imgToInsert)
    except Exception as e:
        logger.error(f"插入图片失败: {Imgpath}, 错误: {str(e)}")


def _escape_uri_path(path: str) -> str:
    """对齐 django.utils.encoding.escape_uri_path（保留 '/'，其余非安全字符转义）"""
    from urllib.parse import quote
    safe = "/:@&=+$,;~()*!'"
    return quote(path, safe=safe)


def get_export_filename(viewset, queryset):
    """取导出文件名（对齐 Django model._meta.verbose_name：
    导出菜单表数据.xlsx / import_template_用户表.xlsx）"""
    from utils.models import _MODEL_VERBOSE_NAMES
    model = getattr(queryset, "model", None) if not isinstance(queryset, type) else queryset
    if hasattr(viewset, "model") and viewset.model:
        model = viewset.model
    if model:
        return (getattr(model, "__export_name__", None)
                or _MODEL_VERBOSE_NAMES.get(getattr(model, "__name__", ""))
                or getattr(model, "__tablename__", None)
                or viewset.file_name)
    return viewset.file_name


def _build_workbook(viewset, data) -> BytesIO | None:
    """按 export_field_dict 生成 Excel（结构对齐 export_data）"""
    wb = Workbook()
    ws = wb.active
    left_center = Alignment(horizontal='left', vertical='center')

    field_order = ["#"] + list(viewset.export_field_dict.keys())
    header_data = ["序号"] + list(viewset.export_field_dict.values())

    df_len_max = [get_string_len(i, viewset.export_column_width) for i in header_data]
    row = get_column_letter(len(viewset.export_field_dict) + 1)
    column = 1

    ws.append(header_data)

    for index, item in enumerate(data):
        row_data = [index + 1]
        for field in field_order[1:]:
            if '.' in field:  # 嵌套字段 c.d
                parts = field.split('.')
                value = item
                try:
                    for part in parts:
                        value = value[part]
                    if isinstance(value, (list, tuple)):
                        value = ",".join(str(v) for v in value)
                    row_data.append(str(value) if value is not None else "")
                except (KeyError, TypeError):
                    row_data.append("")
            else:
                value = item.get(field, "")
                if isinstance(value, (list, tuple)):
                    value = ", ".join(str(v) for v in value)
                row_data.append(str(value) if value is not None else "")

        ws.append(row_data)

        # 特殊字段处理（图片列插入图片，普通列对齐+列宽）
        for h_index, h_item in enumerate(field_order):
            if h_index == 0:
                ws.cell(row=index + 2, column=h_index + 1).alignment = left_center
            else:
                field_name = h_item
                cell_value = row_data[h_index]
                if (field_name in viewset.import_image_fields or
                        ('.' in field_name and field_name.split('.')[-1] in viewset.import_image_fields)):
                    if cell_value and is_image(cell_value):
                        try:
                            inserImg(ws, cell_value, index + 2, h_index + 1)
                            if 10 > df_len_max[h_index]:
                                df_len_max[h_index] = 10
                        except Exception:
                            pass
                else:
                    val = cell_value
                    if isinstance(val, list) and not val:
                        val = ""
                    if isinstance(val, (dict, list)):
                        val = str(val)
                    ws.cell(row=index + 2, column=h_index + 1).value = val
                    ws.cell(row=index + 2, column=h_index + 1).alignment = left_center
                    result_column_width = get_string_len(val, viewset.export_column_width)
                    if h_index != 0 and result_column_width > df_len_max[h_index]:
                        df_len_max[h_index] = result_column_width
        column += 1

    # 更新列宽
    for index, width in enumerate(df_len_max):
        ws.column_dimensions[get_column_letter(index + 1)].width = width

    # 添加过滤头
    from openpyxl.worksheet.table import Table, TableStyleInfo
    tab = Table(displayName="Table", ref=f"A1:{row}{column}")
    style = TableStyleInfo(name="TableStyleLight9", showFirstColumn=True, showLastColumn=True, showRowStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)

    output_buffer = BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer


def export_data_action(viewset):
    """
    导出数据为Excel（对齐 ImportExportMixin.export_data）
    @param: export_fields 前端自定义导出字段；selected_ids 勾选数据id列表
    """
    request_data = viewset.get_request_data()
    export_fields = request_data.get('export_fields', {})
    if export_fields:
        viewset.export_field_dict = export_fields
    selected_ids = request_data.get('selected_ids', [])
    viewset.action = "export_data"
    query = viewset.filter_queryset(viewset.get_queryset())
    if selected_ids:
        query = query.where(viewset.model.id.in_(selected_ids))
    serializer = viewset.get_serializer(viewset._output_serializer())
    items = viewset.db.execute(query).scalars().all()
    data = serializer.serialize_many(items)

    if not data:
        return ErrorResponse(msg='没有可导出的数据')

    try:
        output_buffer = _build_workbook(viewset, data)
        export_file_name = f'导出{get_export_filename(viewset, query)}数据.xlsx'
        return FastAPIResponse(
            content=output_buffer.getvalue(),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={_escape_uri_path(export_file_name)}'},
        )
    except Exception as e:
        logger.error(f"导出失败: {str(e)}")
        return ErrorResponse(msg=f'导出失败: {str(e)}')


def process_nested_field(field_name, value, data_dict, is_many=False, separator=','):
    """处理嵌套字段和多对多字段（对齐 ImportExportMixin.process_nested_field）"""
    if not field_name:
        return
    if value is None:
        value = ""
    value = str(value).strip()

    if is_many:
        related_ids = []
        if value:
            values = [v.strip() for v in value.split(separator) if v.strip()]
            related_ids = values
        if '.' in field_name:
            parts = field_name.split('.')
            current = data_dict
            for part in parts[:-1]:
                if part not in current:
                    current[part] = {}
                current = current[part]
            current[parts[-1]] = related_ids
        else:
            data_dict[field_name] = related_ids
        return

    if '.' in field_name:
        parts = field_name.split('.')
        current = data_dict
        for part in parts[:-1]:
            if part not in current:
                current[part] = {}
            current = current[part]
        current[parts[-1]] = value
    else:
        data_dict[field_name] = value


def import_data_action(viewset, upload_state):
    """批量导入Excel数据（对齐 ImportExportMixin.import_data，事务由会话统一管理）"""
    assert viewset.import_field_dict, "'%s' 请配置对应的导入模板字段。" % viewset.__class__.__name__
    assert viewset.import_serializer_class, "'%s' 请配置对应的导入序列化器。" % viewset.__class__.__name__

    file_info = (upload_state.get("files") or {}).get('file')
    if not file_info:
        return ErrorResponse(msg='请上传文件')

    file_name = file_info.get("name", "")
    if not file_name.endswith(('.xlsx', '.xls')):
        return ErrorResponse(msg='仅支持Excel文件')

    try:
        wb = load_workbook(filename=BytesIO(file_info["content"]))
        ws = wb.active

        headers = [cell.value for cell in ws[1]]

        missing_headers = set(viewset.import_field_dict.keys()) - set(headers)
        if missing_headers:
            return ErrorResponse(msg=f'缺少必要的列: {", ".join(str(m) for m in missing_headers)}')

        valid_data = []
        error_count = 0
        error_messages = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            try:
                row_data = {}
                for idx, cell_value in enumerate(row):
                    if idx >= len(headers):
                        continue
                    excel_col_name = headers[idx]
                    if excel_col_name in viewset.import_field_dict:
                        field_info = viewset.import_field_dict[excel_col_name]
                        model_field = field_info if isinstance(field_info, str) else field_info['field']

                        is_many = False
                        separator = ","
                        if isinstance(field_info, dict):
                            is_many = field_info.get('many', False)
                            separator = field_info.get('separator', ',')

                        if (model_field in viewset.import_image_fields or
                                ('.' in model_field and model_field.split('.')[-1] in viewset.import_image_fields)):
                            cell_value = cell_value if (cell_value and is_image(str(cell_value))) else None

                        cell_value = "" if cell_value is None else cell_value
                        process_nested_field(model_field, cell_value, row_data, is_many, separator)

                valid_data.append(row_data)
            except Exception as e:
                error_count += 1
                error_messages.append(f"第{row_idx}行错误: {str(e)}")
                continue

        if valid_data:
            serializer = viewset.get_serializer(viewset.import_serializer_class)
            created = []
            for row_data in valid_data:
                validated = serializer.run_validation(row_data)
                created.append(serializer.create(validated))
            viewset.db.commit()

        return DetailResponse(data={
            'success': True,
            'success_count': len(valid_data),
            'error_count': error_count,
            'error_messages': error_messages
        })

    except Exception as e:
        logger.error(f"导入失败: {str(e)}", exc_info=True)
        viewset.db.rollback()
        return ErrorResponse(msg=f'导入失败: {str(e)}')


def generate_template(viewset) -> BytesIO:
    """生成简化版导入模板（仅包含字段标题和示例数据）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    left_center = Alignment(horizontal='left', vertical='center')

    headers = []
    example_values = []
    example_data = viewset.get_import_example_data() if hasattr(viewset, 'get_import_example_data') else {}

    for excel_col, field_info in viewset.import_field_dict.items():
        if isinstance(field_info, dict):
            model_field = field_info.get('field', field_info)
            example = field_info.get('example', example_data.get(model_field, ''))
        else:
            example = example_data.get(field_info, '')
        headers.append(excel_col)
        example_values.append(example)

    ws.append(headers)
    ws.append(example_values)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = left_center

    for col in range(1, len(headers) + 1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 20

    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def download_template_action(viewset):
    """下载导入模板（对齐 ImportExportMixin.download_template）"""
    assert viewset.import_field_dict, "'%s' 请配置对应的导入模板字段。" % viewset.__class__.__name__
    try:
        template_stream = generate_template(viewset)
        mname = get_export_filename(viewset, viewset.get_queryset())
        filename = f"import_template_{mname}.xlsx" if mname else f"import_template_{time.strftime('%Y%m%d')}.xlsx"
        return FastAPIResponse(
            content=template_stream.getvalue(),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={_escape_uri_path(filename)}'},
        )
    except Exception as e:
        logger.error(f"生成模板失败: {str(e)}")
        return ErrorResponse(msg=f'生成模板失败: {str(e)}')
